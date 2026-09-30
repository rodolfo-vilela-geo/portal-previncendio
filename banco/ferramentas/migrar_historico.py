"""
Migração da Tabela Principal do BDG (planilha) para a tabela `roi` do Supabase.

Entrada : Tabela_Principal_Banco_de_Dados_Georreferenciado_-_Previncêndio_16.xlsx
Saídas  : roi_historico.csv           -> importar na tabela roi (Table Editor > Import data)
          02_listas_suspensas.sql     -> listas oficiais (ativas) + valores só do histórico (inativos)
          relatorio_migracao.xlsx     -> tudo o que foi corrigido ou esvaziado, registro a registro

Uso: python3 migrar_historico.py caminho/da/planilha.xlsx
"""
import sys, re, csv, datetime as dt, warnings, collections
import openpyxl

warnings.filterwarnings("ignore")
ORIGEM = sys.argv[1]
IMPORTADO_EM = dt.date.today().isoformat()

wb = openpyxl.load_workbook(ORIGEM, read_only=True, data_only=True)
rows = list(wb["Atualizar BDG-2013-2021"].iter_rows(values_only=True))
hdr = rows[0]
H = {h: i for i, h in enumerate(hdr) if h}
registros = [r for r in rows[1:] if any(v is not None for v in r[:83])]

log = []  # (cod_bdp, campo, valor_original, valor_importado, acao)


def registra(cod, campo, orig, novo, acao):
    log.append((cod, campo, str(orig), "" if novo is None else str(novo), acao))


VAZIOS = {"", "-", "- -"}


# ------------------------------------------------------------------ datas
def data(cod, campo, v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    s = str(v).strip()
    if s in VAZIOS:
        return None
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            registra(cod, campo, v, None, "data inexistente — esvaziada")
            return None
    # separador errado ou ponto no fim: 17?11?2023, 08/08/2023.
    m = re.fullmatch(r"(\d{1,2})\D(\d{1,2})\D(\d{4})\.?", s)
    if m:
        d = dt.date(int(m[3]), int(m[2]), int(m[1]))
        registra(cod, campo, v, d, "separador corrigido")
        return d
    # ano com 3 dígitos: 18/08/024
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{3})", s)
    if m:
        d = dt.date(2000 + int(m[3]) % 100, int(m[2]), int(m[1]))
        registra(cod, campo, v, d, "ano corrigido (3 dígitos) — conferir")
        return d
    registra(cod, campo, v, None, "formato não reconhecido — esvaziada")
    return None


def data_aaaammdd(cod, v):
    """data_atua: 20190624 (texto/número), 31102021 (dia-mês-ano), data do Excel ou 'ok'."""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    s = str(v).strip().split(".")[0]
    if re.fullmatch(r"\d{8}", s):
        for a, m_, d_ in ((s[:4], s[4:6], s[6:]), (s[4:], s[2:4], s[:2])):  # AAAAMMDD, DDMMAAAA
            try:
                return dt.date(int(a), int(m_), int(d_))
            except ValueError:
                pass
    if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", s):
        return data(cod, "data_atua", s)
    return None  # 'ok' e similares: sem data


# ------------------------------------------------------------------ horas
def hora(cod, campo, v):
    if v is None:
        return None
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime):
        return v.time()
    s = str(v).strip()
    if s in VAZIOS:
        return None
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", s)
    if m:
        try:
            return dt.time(int(m[1]), int(m[2]), int(m[3] or 0))
        except ValueError:
            registra(cod, campo, v, None, "hora inexistente — esvaziada")
            return None
    m = re.fullmatch(r"(\d{1,2}):0(\d{2}):(\d{2})", s)  # 09:040:00
    if m:
        t = dt.time(int(m[1]), int(m[2]), int(m[3]))
        registra(cod, campo, v, t, "zero extra removido — conferir")
        return t
    registra(cod, campo, v, None, "formato não reconhecido — esvaziada")
    return None


# ------------------------------------------------------------------ números
def numero(cod, campo, v, inteiro=False):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if inteiro:
            if float(v).is_integer():
                return int(v)
            registra(cod, campo, v, round(v), "arredondado para inteiro")
            return round(v)
        return round(float(v), 2)
    s = str(v).strip()
    if s in VAZIOS:
        return None
    # 1,063,00 -> 1063.00 ; 1,10 h -> 1.10
    limpo = re.sub(r"[^\d,\.]", "", s)
    partes = limpo.split(",")
    try:
        if len(partes) >= 2:
            n = float("".join(partes[:-1]) + "." + partes[-1])
        else:
            n = float(limpo)
    except ValueError:
        registra(cod, campo, v, None, "número não reconhecido — esvaziado")
        return None
    n = int(n) if inteiro else round(n, 2)
    registra(cod, campo, v, n, "número convertido")
    return n


def texto(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s in VAZIOS else s


SIM_NAO = {"sim": "Sim", "não": "Não", "nao": "Não"}

INTEIROS = ["comb_uc", "comb_ftp", "comb_sm", "comb_par", "comb_vol", "comb_bm", "comb_pm",
            "vc_4x4", "vc_4x4_d", "vc_4x2", "vc_4x2_d", "vc_van", "vc_van_d", "vc_moto",
            "vc_moto_d", "vc_trator", "vc_trator_d", "vc_pipa", "vc_pipa_d", "a_air_tr", "a_helicop"]
DECIMAIS = ["veg_f_e_d", "veg_f_e_s", "veg_f_o", "veg_c_r", "veg_c_c", "veg_c_a", "veg_c_s_s",
            "veg_c_d", "veg_ver", "veg_ant", "area_int", "area_ent"]
TEXTOS = ["local", "base_op", "ufbio", "categoria", "nome_uc", "grupo", "municipio", "bioma_uc",
          "f_detec", "vc_outro", "veg_outro", "causa_p", "causa_out", "ag_causal", "ag_out",
          "fn_nome", "fn_qnt", "fn_coord", "obs"]

COLUNAS = (["cod_bdp"] + TEXTOS[:9] + ["dat_detec", "hr_detec", "dat_comb", "hr_comb",
           "dat_final", "hr_final"] + INTEIROS + ["vc_outro"] + DECIMAIS + ["veg_outro",
           "causa_p", "causa_out", "ag_causal", "ag_out", "ind_aut", "fn_nome", "fn_qnt",
           "fn_coord", "obs", "status", "enviado_em", "atualizado_em"])

saida = []
sem_data_atua = 0
for r in registros:
    g = lambda c: r[H[c]]
    cod = str(g("cod_bdp")).strip()
    reg = {"cod_bdp": cod}
    for c in TEXTOS:
        reg[c] = texto(g(c))
    for c in ("dat_detec", "dat_comb", "dat_final"):
        reg[c] = data(cod, c, g(c))
    for c in ("hr_detec", "hr_comb", "hr_final"):
        reg[c] = hora(cod, c, g(c))
    for c in INTEIROS:
        reg[c] = numero(cod, c, g(c), inteiro=True)
    for c in DECIMAIS:
        reg[c] = numero(cod, c, g(c))
    ia = texto(g("ind_aut"))
    if ia is not None:
        norm = SIM_NAO.get(ia.lower())
        if norm != ia:
            registra(cod, "ind_aut", ia, norm, "padronizado" if norm else "valor fora da lista — esvaziado")
        ia = norm
    reg["ind_aut"] = ia
    reg["status"] = "historico"
    reg["enviado_em"] = None
    da = data_aaaammdd(cod, g("data_atua"))
    if da is None:
        sem_data_atua += 1
        da = IMPORTADO_EM
    reg["atualizado_em"] = da
    saida.append(reg)

# ------------------------------------------------------------------ CSV
with open("roi_historico.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(COLUNAS)
    for reg in saida:
        w.writerow(["" if reg[c] is None else (reg[c].isoformat() if hasattr(reg[c], "isoformat") else reg[c])
                    for c in COLUNAS])

# ------------------------------------------------------------------ listas suspensas
LISTAS = {"Local": "local", "BASE_FTP": "base_op", "REG_IEF": "ufbio", "CATEGORIA": "categoria",
          "NOME_UC": "nome_uc", "GRUPO": "grupo", "MUNICIPIOS": "municipio", "BIOMA": "bioma_uc",
          "F_DETEC": "f_detec", "CAU_PRINC": "causa_p", "AG_CAU": "ag_causal", "IND_AU": "ind_aut"}
lrows = list(wb["LISTAS SUSPENSAS"].iter_rows(values_only=True))
oficiais = collections.defaultdict(list)
for j, h in enumerate(lrows[0]):
    if h in LISTAS:
        for r in lrows[1:]:
            v = texto(r[j]) if j < len(r) else None
            if v and v not in oficiais[LISTAS[h]]:
                oficiais[LISTAS[h]].append(v)
linhas, inativos = [], collections.Counter()
q = lambda s: "'" + s.replace("'", "''") + "'"
for campo, vals in oficiais.items():
    for k, v in enumerate(vals, 1):
        linhas.append(f"  ({q(campo)}, {q(v)}, {k}, true)")
    usados = sorted({reg[campo] for reg in saida if reg[campo]} - set(vals))
    for v in usados:
        linhas.append(f"  ({q(campo)}, {q(v)}, 999, false)")
        inativos[campo] += 1
with open("02_listas_suspensas.sql", "w", encoding="utf-8") as f:
    f.write("-- Listas suspensas (aba LISTAS SUSPENSAS da Tabela Principal BDG v16).\n"
            "-- ativo = true  : valores oficiais, oferecidos no formulário.\n"
            "-- ativo = false : valores que só aparecem no histórico; ficam registrados, mas o\n"
            "--                 formulário não oferece. Revisar: padronizar ou promover a oficial.\n"
            "-- Rodar depois de 01_esquema_roi.sql.\n"
            "insert into dominio (campo, valor, ordem, ativo) values\n")
    f.write(",\n".join(linhas) + "\non conflict do nothing;\n")

# ------------------------------------------------------------------ relatório
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
rel = Workbook()
ws = rel.active
ws.title = "Correções"
ws.append(["cod_bdp", "campo", "valor na planilha", "valor importado", "ação"])
for linha in sorted(log):
    ws.append(list(linha))
ws2 = rel.create_sheet("Fora das listas")
ws2.append(["campo", "valor usado no histórico", "nº de registros"])
for campo, vals in oficiais.items():
    cont = collections.Counter(reg[campo] for reg in saida if reg[campo] and reg[campo] not in vals)
    for v, n in cont.most_common():
        ws2.append([campo, v, n])
for s in (ws, ws2):
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7A1F1F")
    s.freeze_panes = "A2"
    for col, larg in zip("ABCDE", (16, 14, 40, 24, 48)):
        s.column_dimensions[col].width = larg
rel.save("relatorio_migracao.xlsx")

print(f"{len(saida)} registros -> roi_historico.csv")
print(f"{len(log)} correções registradas:", dict(collections.Counter(l[4] for l in log)))
print(f"{sem_data_atua} registros sem data_atua válida: recebem a data da importação")
print("valores fora das listas (inativos):", dict(inativos))
